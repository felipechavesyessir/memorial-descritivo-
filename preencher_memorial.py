#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.document import Document as DocumentType
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

DOTTED_PLACEHOLDER_RE = re.compile(r"\.{5,}")


def extract_pdf_text(pdf_path: Path) -> str:
    reader = PdfReader(str(pdf_path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip(" .;:\n\t")


def parse_confrontantes(text: str) -> Dict[str, str]:
    normalized = re.sub(r"\s+", " ", text)
    labels = {
        "norte": ["NORTE"],
        "sul": ["SUL"],
        "leste": ["LESTE", "NASCENTE"],
        "oeste": ["OESTE", "POENTE"],
    }

    found: Dict[str, str] = {}

    all_tokens = [token for values in labels.values() for token in values]
    next_label_pattern = r"(?:" + "|".join(all_tokens) + r")"

    for side, tokens in labels.items():
        for token in tokens:
            pattern = re.compile(
                rf"\b{token}\b\s*[:\-–]\s*(.+?)(?=\b{next_label_pattern}\b\s*[:\-–]|$)",
                re.IGNORECASE,
            )
            match = pattern.search(normalized)
            if match:
                found[side] = _normalize(match.group(1))
                break

    return found


def _iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from _iter_cell_paragraphs(cell)


def iter_document_paragraphs(doc: DocumentType) -> Iterable[Paragraph]:
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)


def replace_first_placeholder_in_paragraph(paragraph: Paragraph, value: str) -> bool:
    if not paragraph.runs:
        return False

    full_text = "".join(run.text for run in paragraph.runs)
    match = DOTTED_PLACEHOLDER_RE.search(full_text)
    if not match:
        return False

    start, end = match.span()
    run_ranges = []
    index = 0
    for i, run in enumerate(paragraph.runs):
        next_index = index + len(run.text)
        run_ranges.append((i, index, next_index))
        index = next_index

    start_info = next((r for r in run_ranges if r[1] <= start < r[2]), None)
    end_info = next((r for r in run_ranges if r[1] < end <= r[2]), None)
    if start_info is None or end_info is None:
        return False

    start_run, start_run_begin, _ = start_info
    end_run, end_run_begin, _ = end_info

    prefix = paragraph.runs[start_run].text[: start - start_run_begin]
    suffix = paragraph.runs[end_run].text[end - end_run_begin :]

    paragraph.runs[start_run].text = prefix + value

    for i in range(start_run + 1, end_run):
        paragraph.runs[i].text = ""

    if end_run != start_run:
        paragraph.runs[end_run].text = suffix
    else:
        paragraph.runs[start_run].text = prefix + value + suffix

    return True


def fill_docx_placeholders(template_path: Path, output_path: Path, replacements: List[str]) -> int:
    doc = Document(str(template_path))
    queue = list(replacements)
    replaced = 0

    for paragraph in iter_document_paragraphs(doc):
        while queue and replace_first_placeholder_in_paragraph(paragraph, queue[0]):
            queue.pop(0)
            replaced += 1

    doc.save(str(output_path))
    return replaced


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Preenche um memorial descritivo DOCX com confrontantes extraídos de um PDF."
    )
    parser.add_argument("--pdf", required=True, type=Path, help="Arquivo PDF com confrontantes")
    parser.add_argument("--template", required=True, type=Path, help="Template DOCX com '..........'")
    parser.add_argument("--output", required=True, type=Path, help="Arquivo DOCX de saída")

    args = parser.parse_args()

    pdf_text = extract_pdf_text(args.pdf)
    confrontantes = parse_confrontantes(pdf_text)

    order = ["norte", "sul", "leste", "oeste"]
    values = [confrontantes.get(direction, "") for direction in order]

    if not any(values):
        raise SystemExit(
            "Não foi possível identificar confrontantes no PDF. Verifique se ele contém rótulos como NORTE/SUL/LESTE/OESTE."
        )

    replaced = fill_docx_placeholders(args.template, args.output, values)
    print(f"Confrontantes extraídos: {confrontantes}")
    print(f"Placeholders substituídos: {replaced}")
    print(f"Arquivo gerado: {args.output}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
