import tempfile
import unittest
from pathlib import Path

from docx import Document

from preencher_memorial import (
    fill_docx_placeholders,
    parse_confrontantes,
    replace_first_placeholder_in_paragraph,
)


class TestParseConfrontantes(unittest.TestCase):
    def test_parse_with_cardinal_points(self):
        text = (
            "NORTE: João da Silva; SUL: Estrada Municipal; "
            "LESTE: Rio Alegre; OESTE: Fazenda Boa Vista"
        )
        parsed = parse_confrontantes(text)

        self.assertEqual(parsed["norte"], "João da Silva")
        self.assertEqual(parsed["sul"], "Estrada Municipal")
        self.assertEqual(parsed["leste"], "Rio Alegre")
        self.assertEqual(parsed["oeste"], "Fazenda Boa Vista")

    def test_parse_with_nascente_poente(self):
        text = "NORTE: Area A SUL: Area B NASCENTE: Rua C POENTE: Sitio D"
        parsed = parse_confrontantes(text)

        self.assertEqual(parsed["leste"], "Rua C")
        self.assertEqual(parsed["oeste"], "Sitio D")


class TestPlaceholderReplacement(unittest.TestCase):
    def test_replace_first_placeholder_in_paragraph(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Confrontando com ")
        paragraph.add_run("..........").bold = True
        paragraph.add_run(" até o marco.")

        replaced = replace_first_placeholder_in_paragraph(paragraph, "João")

        self.assertTrue(replaced)
        self.assertEqual(paragraph.text, "Confrontando com João até o marco.")

    def test_fill_docx_placeholders(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            template = tmp / "template.docx"
            output = tmp / "output.docx"

            doc = Document()
            doc.add_paragraph("Norte: ..........")
            doc.add_paragraph("Sul: ..........")
            doc.add_paragraph("Leste: ..........")
            doc.add_paragraph("Oeste: ..........")
            doc.save(template)

            replaced = fill_docx_placeholders(
                template,
                output,
                ["A", "B", "C", "D"],
            )

            self.assertEqual(replaced, 4)

            generated = Document(output)
            text = "\n".join(p.text for p in generated.paragraphs)
            self.assertIn("Norte: A", text)
            self.assertIn("Sul: B", text)
            self.assertIn("Leste: C", text)
            self.assertIn("Oeste: D", text)


if __name__ == "__main__":
    unittest.main()
